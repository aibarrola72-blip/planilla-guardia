"""Tests del núcleo de importación desde Excel (sin red)."""

from __future__ import annotations

import io
from datetime import date

from openpyxl import Workbook, load_workbook

from app import importar_personas as ie


def cat_base() -> ie.Catalogos:
    return ie.Catalogos(
        unidades={"urgencias": 1, "pediatria": 2},
        nombres_unidad={1: "URGENCIAS", 2: "PEDIATRIA"},
        sectores={(1, "rac"): 10},
        cargos={"rt": 100},
        turnos={"M": 200, "T": 201, "N1": 202},
    )


def fila(**kw) -> ie.FilaPersonal:
    base = {"linea": 2, "nombre": "Lic. Ana Torres", "unidad": "URGENCIAS",
            "sector": "RAC", "cargo": "RT"}
    base.update(kw)
    return ie.FilaPersonal(**base)


def bd_base() -> list[dict]:
    return [
        {"id": 1, "nombre": "Lic. Ana Torres", "ci": "111", "registro": "R1",
         "unidad_id": 1, "sector_id": 10, "cargo_id": 100, "turno_id": 200,
         "estado": "ACTIVO", "orden": 1, "idpersonal_legacy": None},
        {"id": 2, "nombre": "Lic. Juan Lopez", "ci": "222", "registro": "R2",
         "unidad_id": 1, "sector_id": 10, "cargo_id": 100, "turno_id": 202,
         "estado": "ACTIVO", "orden": 2, "idpersonal_legacy": None},
        {"id": 3, "nombre": "Dr. De Otra Uni", "ci": "333", "registro": None,
         "unidad_id": 9, "sector_id": None, "cargo_id": None, "turno_id": None,
         "estado": "ACTIVO", "orden": 1, "idpersonal_legacy": None},
        {"id": 4, "nombre": "Lic. Baja Ya", "ci": "444", "registro": None,
         "unidad_id": 1, "sector_id": 10, "cargo_id": 100, "turno_id": 201,
         "estado": "INACTIVO", "orden": 3, "idpersonal_legacy": None},
    ]


def test_leer_nomina() -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "NOMINA"
    ws.append(["id_NOMINA", "PERSONAL", "UNIDAD", "SECTOR", "CARGO", "TURNO",
               "C.I. N\u00b0", "Reg. N\u00b0", "ESTADO"])
    ws.append(["A1", "Lic. Ana Torres", "URGENCIAS", "RAC", "RT", "1",
               11111111.0, "R1", "ACTIVO"])
    ws.append([None, "", "", "", "", "", "", "", ""])
    ws.append([None, "Lic. Nuevo", "URGENCIAS", "RAC", "RT", "M", 22222222, "", ""])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    filas = ie.leer_nomina(ie.leer_hoja(load_workbook(buf, data_only=True), "NOMINA"))
    assert len(filas) == 2, filas
    assert filas[0].ci == "11111111" and filas[0].id_legacy == "A1"
    assert filas[0].turno == "1" and filas[0].estado == "ACTIVO"
    assert filas[1].turno == "M" and filas[1].estado == "ACTIVO"
    print("OK leer_nomina: encabezados, CI numerica y filas vacias")


def test_hojas_invalidas() -> None:
    wb = Workbook()
    wb.active.title = "OTRA"
    try:
        ie.leer_hoja(wb, "NOMINA")
        raise AssertionError("debia fallar hoja inexistente")
    except ie.ErrorImportacion as exc:
        assert "NOMINA" in str(exc)
    print("OK ErrorImportacion ante hoja inexistente")


def test_alta_y_match() -> None:
    filas = [fila(), fila(linea=3, nombre="Lic. Nuevo", ci="777")]
    cambios = ie.calcular_cambios(filas, bd_base(), cat_base())
    assert len(cambios.altas) == 1, cambios.altas
    assert cambios.altas[0]["nombre"] == "Lic. Nuevo"
    assert cambios.altas[0]["unidad_id"] == 1 and cambios.altas[0]["orden"] == 4
    assert not cambios.revision, cambios.revision
    assert not cambios.mods, cambios.mods
    print("OK alta nueva + match por CI sin diff")


def test_mod_y_campo_vacio_no_borra() -> None:
    filas = [fila(turno="T", registro="")]  # registro vacio: no debe borrar R1
    cambios = ie.calcular_cambios(filas, bd_base(), cat_base())
    assert cambios.mods and cambios.mods[0][0] == 1, cambios.mods
    diff = cambios.mods[0][1]
    assert diff == {"turno_id": 201}, diff
    assert "registro" not in diff and "ci" not in diff
    print("OK mod por turno; campo vacio no borra datos")


def test_bajas_acotadas() -> None:
    filas = [fila()]  # solo unidad 1 presente; matchea a la id 1
    cambios = ie.calcular_cambios(filas, bd_base(), cat_base())
    assert [b[0] for b in cambios.bajas] == [2], cambios.bajas
    global_ = ie.calcular_cambios(filas, bd_base(), cat_base(), bajas_global=True)
    assert sorted(b[0] for b in global_.bajas) == [2, 3], global_.bajas
    print("OK bajas acotadas a las unidades del archivo (+ bajas_global)")


def test_revisiones() -> None:
    bd = bd_base() + [
        {"id": 5, "nombre": "Dup", "ci": "999", "unidad_id": 1, "estado": "ACTIVO",
         "orden": 9, "sector_id": None, "cargo_id": None, "turno_id": None,
         "registro": None, "idpersonal_legacy": None},
        {"id": 6, "nombre": "Dup2", "ci": "999", "unidad_id": 1, "estado": "ACTIVO",
         "orden": 10, "sector_id": None, "cargo_id": None, "turno_id": None,
         "registro": None, "idpersonal_legacy": None},
    ]
    filas = [
        fila(linea=2, nombre="Dup", ci="999"),                    # CI duplicada en BD
        fila(linea=3, nombre="Lic. Sin Uni", unidad="NOEXISTE"),  # catalogo faltante
        fila(linea=4, nombre="Lic. Ana Torres"),                  # misma persona 2 veces
        fila(linea=5, nombre="Lic. Ana Torres"),
    ]
    cambios = ie.calcular_cambios(filas, bd, cat_base())
    assert not cambios.altas and not cambios.mods, (cambios.altas, cambios.mods)
    assert len(cambios.revision) == 3, cambios.revision
    assert any("duplicada en BD" in r for r in cambios.revision)
    assert any("unidad 'NOEXISTE'" in r for r in cambios.revision)
    assert any("repetida en el archivo" in r for r in cambios.revision)
    print("OK revisiones: CI duplicada, catalogo faltante y fila repetida")


def test_ausencias_dedupe() -> None:
    bd = bd_base()
    clave = (1, "2026-10-01", "2026-10-05")
    vac = [("111", "", date(2026, 10, 1), date(2026, 10, 5), "")]
    out = ie.calcular_ausencias(vac, [], bd, {clave}, set())
    assert out.vacaciones == [], out.vacaciones          # ya existe -> no duplica

    out = ie.calcular_ausencias(
        [("111", "", date(2026, 11, 1), date(2026, 11, 3), "")], [], bd, set(), set())
    assert len(out.vacaciones) == 1 and out.vacaciones[0]["persona_id"] == 1

    libres = [("111", "", date(2026, 12, 1), None, ""),
              ("111", "", date(2026, 12, 2), None, "")]
    out = ie.calcular_ausencias([], libres, bd, set(), set())
    assert len(out.libres) == 1 and out.libres[0]["fecha_fin"] == "2026-12-02"

    out = ie.calcular_ausencias(
        [("9999", "Nadie", date(2026, 1, 1), date(2026, 1, 2), "")], [], bd, set(), set())
    assert out.omitidas and not out.vacaciones
    print("OK ausencias: dedupe, rango de libres y filas sin persona")


def test_informe_dict() -> None:
    filas = [fila(turno="T")]
    cambios = ie.calcular_cambios(filas, bd_base(), cat_base())
    inf = ie.Informe(filas=len(filas), cambios=cambios, ausencias=None,
                     catalogos={"unidades": [], "sectores": [], "cargos": []},
                     aplicado=False, nombres={1: "Lic. Ana Torres"})
    d = inf.como_dict()
    assert d["aplicado"] is False and d["filas"] == 1
    assert d["mods"][0]["nombre"] == "Lic. Ana Torres"
    assert d["bajas"][0]["nombre"] == "Lic. Juan Lopez"
    assert "altas" in d and "revision" in d and "ausencias" in d
    print("OK contrato JSON del endpoint (Informe.como_dict)")


if __name__ == "__main__":
    test_leer_nomina()
    test_hojas_invalidas()
    test_alta_y_match()
    test_mod_y_campo_vacio_no_borra()
    test_bajas_acotadas()
    test_revisiones()
    test_ausencias_dedupe()
    test_informe_dict()
    print("\ntodos los tests OK")
